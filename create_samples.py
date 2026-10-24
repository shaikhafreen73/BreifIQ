import os
import pandas as pd
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(SAMPLE_DIR, exist_ok=True)


def create_sample_pdf():
    pdf_path = os.path.join(SAMPLE_DIR, "sample_ai_research.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor('#1e1b4b'))
    head_style = ParagraphStyle('Heading', parent=styles['Heading2'], fontSize=13, leading=16, textColor=colors.HexColor('#4338ca'), spaceBefore=10)
    body = ParagraphStyle('Body', parent=styles['BodyText'], fontSize=10, leading=14, textColor=colors.HexColor('#334155'))

    story = [
        Paragraph("Advances in Sequence-to-Sequence Deep Learning for Text Summarization", title_style),
        Paragraph("Dr. Elena Vance and Dr. Marcus Thorne • AI Research Consortium", ParagraphStyle('Sub', parent=styles['Normal'], fontSize=9, textColor=colors.gray)),
        Spacer(1, 12),
        Paragraph("Abstract", head_style),
        Paragraph(
            "Natural Language Processing (NLP) has undergone a fundamental transformation with the arrival of Transformer architectures. "
            "Text summarization represents one of the most intellectually demanding challenges in artificial intelligence, "
            "requiring models to simultaneously comprehend contextual semantics, synthesize disparate conceptual ideas, "
            "and distill verbose passages into coherent, concise executive briefings. "
            "The Text-to-Text Transfer Transformer (T5) re-conceptualizes all NLP tasks into a unified sequence-to-sequence format, "
            "leveraging pre-trained masked language modeling to achieve superior abstractive summarization performance.",
            body
        ),
        Paragraph("Key Methodological Innovations", head_style),
        Paragraph(
            "1. Bidirectional Attention Mechanisms: Unlike unidirectional recurrent architectures, transformer self-attention enables token-level "
            "correlation across arbitrary distances within textual corpora, eliminating the vanishing gradient dilemma during document analysis.\n\n"
            "2. Abstractive versus Extractive Syntheses: While legacy systems merely lifted verbatim sentences from source documents, "
            "deep learning seq2seq frameworks generate novel lexical expressions and paraphrasing that emulate human cognitive synthesis.\n\n"
            "3. Computational Scalability: Utilizing tokenized sub-word embeddings (SentencePiece) and parameter-efficient fine-tuning, "
            "compact models like T5-small deliver exceptional inference latency while maintaining high ROUGE benchmark fidelity.",
            body
        ),
        Paragraph("Empirical Findings & Real-World Utility", head_style),
        Paragraph(
            "Our experimental benchmarks indicate that automated summarization systems reduce academic and enterprise document review times "
            "by over 78%. By translating lengthy research publications and tabular reports into digestible executive digests, "
            "students, educators, and enterprise analysts can rapidly identify critical insights, detect emergent thematic patterns, "
            "and execute informed strategic decisions without cognitive fatigue.",
            body
        ),
        Paragraph("Conclusion", head_style),
        Paragraph(
            "The integration of deep learning summarization with automated tabular data reporting establishes a new standard for intelligent knowledge retrieval. "
            "As neural language models continue to evolve, multi-modal ingestion and explainable AI pipelines will further democratize data accessibility.",
            body
        )
    ]

    doc.build(story)
    print("Created sample_ai_research.pdf")


def create_sample_csv():
    csv_path = os.path.join(SAMPLE_DIR, "sample_student_grades.csv")
    data = {
        "Student Name": [
            "Sophia Anderson", "Ethan Brown", "Olivia Davis", "Liam Garcia",
            "Emma Johnson", "Noah Martinez", "Ava Miller", "Lucas Robinson",
            "Isabella Taylor", "Mason Wilson", "Mia White", "James Harris"
        ],
        "Attendance (%)": [94.5, 72.0, 98.0, 81.5, 68.0, 91.0, 88.5, 74.0, 96.0, 62.5, 85.0, 93.0],
        "Overall Score (%)": [92.0, 78.5, 95.0, 84.0, 64.0, 89.5, 88.0, 71.5, 94.0, 58.0, 82.5, 91.5],
        "Performance": [
            "Excellent", "Good", "Excellent", "Very Good", "Average",
            "Very Good", "Very Good", "Good", "Excellent", "Needs Improvement",
            "Very Good", "Excellent"
        ],
        "Assignments Completed": [10, 8, 10, 9, 6, 9, 9, 7, 10, 5, 8, 10]
    }
    df = pd.DataFrame(data)
    df.to_csv(csv_path, index=False)
    print("Created sample_student_grades.csv")


def create_sample_excel():
    excel_path = os.path.join(SAMPLE_DIR, "sample_sales_analytics.xlsx")
    data = {
        "Product": ["Quantum Pro AI", "CloudSync Enterprise", "DataFlow Suite", "CyberShield Plus", "SmartBI Dashboard"],
        "Category": ["Software", "Cloud", "Analytics", "Security", "Analytics"],
        "Units Sold": [1240, 850, 2100, 1680, 1420],
        "Revenue ($)": [186000, 127500, 315000, 201600, 170400],
        "Customer Rating": [4.8, 4.5, 4.9, 4.7, 4.6]
    }
    df = pd.DataFrame(data)
    df.to_excel(excel_path, index=False)
    print("Created sample_sales_analytics.xlsx")


def create_sample_docx():
    docx_path = os.path.join(SAMPLE_DIR, "sample_executive_brief.docx")
    doc = docx.Document()
    doc.add_heading("BriefIQ Strategic Initiative Overview", level=0)
    doc.add_paragraph(
        "BriefIQ is an AI-powered intelligence platform engineered to bridge the gap between extensive unstructured textual "
        "documentation and high-density quantitative tabular datasets. By embedding deep learning transformers directly into "
        "an intuitive web interface, BriefIQ accelerates knowledge discovery for academic institutions and enterprise teams."
    )
    doc.add_heading("Key Operational Objectives", level=1)
    doc.add_paragraph("1. Zero Latency Summarization: Rapid generation of key takeaways from PDF and Word documents.")
    doc.add_paragraph("2. Automated Quantitative EDA: Instant statistical profiling and chart visualization for CSV and Excel files.")
    doc.add_paragraph("3. Actionable Advisory Reports: Personalized feedback and performance diagnosis.")
    doc.save(docx_path)
    print("Created sample_executive_brief.docx")


if __name__ == "__main__":
    create_sample_pdf()
    create_sample_csv()
    create_sample_excel()
    create_sample_docx()
    print("All sample files generated successfully!")
